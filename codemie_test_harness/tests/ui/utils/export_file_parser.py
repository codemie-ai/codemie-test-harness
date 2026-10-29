# Copyright 2026 EPAM Systems, Inc. (“EPAM”)
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Helpers that read files downloaded by the "Export Conversation" dialog.

DOCX is parsed with python-docx, PDF with pypdf and JSON with the stdlib, so tests
can assert on what the exported document actually contains.
"""

import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from codemie_test_harness.tests.ui.test_data.export_conversation_test_data import (
    ExportFormat,
)

MONOSPACE_FONTS = {"consolas", "courier new", "courier", "menlo", "monaco"}


@dataclass
class DocxParagraph:
    """A DOCX paragraph: style name, list kind ("bullet"/"numbered"/None), monospace flag."""

    style: str
    text: str
    list_kind: str | None = None
    monospace: bool = False


@dataclass
class DocxContent:
    """Structural view of a DOCX export."""

    paragraphs: list[DocxParagraph] = field(default_factory=list)
    tables: list[list[list[str]]] = field(default_factory=list)
    hyperlinks: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        """All paragraph and table-cell text joined into one string."""
        cells = [cell for table in self.tables for row in table for cell in row]
        return "\n".join([p.text for p in self.paragraphs] + cells)

    def _texts(self, predicate) -> list[str]:
        return [p.text for p in self.paragraphs if predicate(p)]

    @property
    def heading_1(self) -> list[str]:
        return self._texts(lambda p: p.style == "Heading 1")

    @property
    def heading_2(self) -> list[str]:
        return self._texts(lambda p: p.style == "Heading 2")

    @property
    def bullets(self) -> list[str]:
        return self._texts(lambda p: p.list_kind == "bullet")

    @property
    def numbered(self) -> list[str]:
        return self._texts(lambda p: p.list_kind == "numbered")

    @property
    def monospace_blocks(self) -> list[str]:
        """Text of paragraphs whose every run is rendered in a monospace font."""
        return self._texts(lambda p: p.monospace)


def _list_kind(document, paragraph) -> str | None:
    num_pr = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
    if num_pr is None:
        return None
    ilvl = num_pr.ilvl.val if num_pr.ilvl is not None else 0
    numbering = document.part.numbering_part.element
    abstract_id = numbering.xpath(
        f"w:num[@w:numId='{num_pr.numId.val}']/w:abstractNumId/@w:val"
    )[0]
    num_format = numbering.xpath(
        f"w:abstractNum[@w:abstractNumId='{abstract_id}']"
        f"/w:lvl[@w:ilvl='{ilvl}']/w:numFmt/@w:val"
    )[0]
    return "bullet" if num_format == "bullet" else "numbered"


def _run_font_name(run) -> str | None:
    """Font of a run: direct formatting first, then its character-style chain."""
    if run.font.name:
        return run.font.name
    style = run.style
    while style is not None:
        if style.font.name:
            return style.font.name
        style = style.base_style
    return None


def _is_monospace(paragraph) -> bool:
    runs = [run for run in paragraph.runs if run.text.strip()]
    return bool(runs) and all(
        (_run_font_name(run) or "").lower() in MONOSPACE_FONTS for run in runs
    )


def extract_docx(path: Path) -> DocxContent:
    """Parse a DOCX export into paragraphs (style/list kind/font), tables and hyperlink targets."""
    document = Document(str(path))
    return DocxContent(
        paragraphs=[
            DocxParagraph(
                p.style.name, p.text, _list_kind(document, p), _is_monospace(p)
            )
            for p in document.paragraphs
        ],
        tables=[
            [[cell.text for cell in row.cells] for row in table.rows]
            for table in document.tables
        ],
        hyperlinks=[
            rel.target_ref
            for rel in document.part.rels.values()
            if rel.is_external and rel.reltype.endswith("/hyperlink")
        ],
    )


def extract_pdf_text(path: Path) -> str:
    """Return the text of all PDF pages, with layout-inserted line-end hyphens removed."""
    text = "\n".join(page.extract_text() for page in PdfReader(str(path)).pages)
    return re.sub(r"-\n", "", text)


def extract_pdf_links(path: Path) -> list[str]:
    """Return the URI targets of all link annotations in the PDF (resolving indirect objects)."""
    links = []
    for page in PdfReader(str(path)).pages:
        annotations = page.get("/Annots")
        if annotations is None:
            continue
        for annotation in annotations.get_object():
            action = annotation.get_object().get("/A")
            if action is not None and "/URI" in action.get_object():
                links.append(str(action.get_object()["/URI"]))
    return links


def extract_text(path: Path, export_format: ExportFormat) -> str:
    """Return the readable text of an exported file in any format."""
    if export_format is ExportFormat.DOCX:
        return extract_docx(path).text
    if export_format is ExportFormat.PDF:
        return extract_pdf_text(path)
    return path.read_text(encoding="utf-8")


def load_json_export(path: Path) -> dict:
    """Load a JSON conversation export."""
    return json.loads(path.read_text(encoding="utf-8"))


def json_thought_authors(path: Path) -> list[str]:
    """Return the author names of all thoughts (tool outputs) in a JSON export."""
    return [
        thought["author_name"]
        for message in load_json_export(path)["history"]
        for thought in message.get("thoughts") or []
    ]
